# -*- coding: utf-8 -*-
"""
批量导入模块：把一份表格（Excel / CSV）解析成结构化的知识条目，供 /import 使用。

设计要点：
- 只做"解析 + 差异比对"，**不直接落库**。落库由调用方（commands / kb_engine）在
  用户确认后执行，避免误传文件污染知识库。
- 支持两种来源：
  * 本地路径（服务器上的文件，含共享盘 UNC 路径）
  * http/https 链接（先用 urllib 下载到临时文件再解析）
- 支持 .xlsx（用 openpyxl，已预装）与 .csv（utf-8-sig / gbk 自动尝试）。
- 表头兼容中文列名，缺列自动补空；"标准问题"必填，"标准答案"为空则视为无效行。
"""
import csv
import io
import os
import tempfile
import urllib.request

# 与 kb.csv 对齐的列（"id""处理人""更新时间"由系统生成，导入时不要求）
FIELD_ALIASES = {
    '标准问题': ['标准问题', '问题', '标题', 'question', 'title'],
    '标准答案': ['标准答案', '答案', '内容', 'answer', 'content'],
    '相似问法': ['相似问法', '相似问题', '同义问法', 'simquestions'],
    '触发关键词': ['触发关键词', '关键词', '关键字', 'keywords'],
    '类别': ['类别', '一级目录', '一级', 'category'],
    '分组': ['分组', '二级目录', '二级', 'group'],
    '三级目录': ['三级目录', '三级', '子目录', 'sub'],
    '升级路径': ['升级路径', '升级', 'escalation'],
}

MAX_BYTES = 8 * 1024 * 1024     # 单文件上限 8MB，防止误传大文件


def _norm_header(h):
    return (h or '').strip().lower().replace(' ', '').replace('_', '')


def _map_columns(headers):
    """把实际表头映射成标准字段名。返回 {标准字段: 列下标}"""
    norm = [_norm_header(h) for h in headers]
    mapping = {}
    for std, aliases in FIELD_ALIASES.items():
        for a in aliases:
            key = _norm_header(a)
            if key in norm:
                mapping[std] = norm.index(key)
                break
    return mapping


def _cell(row, mapping, field):
    idx = mapping.get(field)
    if idx is None or idx >= len(row):
        return ''
    v = row[idx]
    if v is None:
        return ''
    return str(v).strip()


def parse_rows(headers, rows):
    """把表头 + 数据行解析为条目列表与问题清单。

    返回 (records, errors)
      records: [{'类别':..,'分组':..,'三级目录':..,'标准问题':..,'标准答案':..,
                 '相似问法':..,'触发关键词':..,'升级路径':..}, ...]
      errors : ['第3行：缺少标准问题', ...]
    """
    mapping = _map_columns(headers)
    if '标准问题' not in mapping:
        return [], ['表头里找不到「标准问题」列（也可写作：问题/标题）']
    if '标准答案' not in mapping:
        return [], ['表头里找不到「标准答案」列（也可写作：答案/内容）']

    records, errors = [], []
    for i, row in enumerate(rows, start=2):     # 第1行是表头
        if not row or all((c is None or str(c).strip() == '') for c in row):
            continue                            # 跳过空行
        q = _cell(row, mapping, '标准问题')
        a = _cell(row, mapping, '标准答案')
        if not q:
            errors.append(f'第{i}行：缺少「标准问题」，已跳过')
            continue
        if not a:
            errors.append(f'第{i}行：「{q}」缺少「标准答案」，已跳过')
            continue
        records.append({
            '类别': _cell(row, mapping, '类别'),
            '分组': _cell(row, mapping, '分组'),
            '三级目录': _cell(row, mapping, '三级目录'),
            '标准问题': q,
            '标准答案': a,
            '相似问法': _cell(row, mapping, '相似问法'),
            '触发关键词': _cell(row, mapping, '触发关键词'),
            '升级路径': _cell(row, mapping, '升级路径'),
        })
    return records, errors


def parse_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        headers = next(it)
    except StopIteration:
        return [], ['表格是空的']
    rows = [list(r) for r in it]
    wb.close()
    return parse_rows(list(headers), rows)


def parse_csv(path):
    for enc in ('utf-8-sig', 'gbk', 'utf-8'):
        try:
            with open(path, encoding=enc, newline='') as f:
                reader = csv.reader(f)
                all_rows = list(reader)
            break
        except UnicodeDecodeError:
            continue
    else:
        return [], ['无法识别文件编码（请另存为 UTF-8 或 GBK 的 CSV）']
    if not all_rows:
        return [], ['文件是空的']
    return parse_rows(all_rows[0], all_rows[1:])


def parse_file(path):
    """按扩展名选择解析器。返回 (records, errors)"""
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        return parse_xlsx(path)
    if ext == '.csv':
        return parse_csv(path)
    if ext == '.xls':
        return [], ['暂不支持老版 .xls，请在 Excel 里另存为 .xlsx 或 .csv']
    return [], [f'不支持的文件类型：{ext or "（无扩展名）"}（仅支持 .xlsx / .csv）']


def is_url(src):
    return src.lower().startswith(('http://', 'https://'))


def download(url):
    """下载链接到临时文件。返回 (本地路径, 错误信息)"""
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'IT-KB-Bot/1.0'})
        with urllib.request.urlopen(req, timeout=20) as resp:
            data = resp.read(MAX_BYTES + 1)
        if len(data) > MAX_BYTES:
            return None, f'文件超过 {MAX_BYTES // 1024 // 1024}MB 上限，请拆分后再导入'
        # 从 URL 里猜扩展名
        ext = os.path.splitext(url.split('?')[0])[1].lower()
        if ext not in ('.xlsx', '.xlsm', '.csv'):
            ext = '.xlsx'
        fd, tmp = tempfile.mkstemp(suffix=ext, prefix='kbimp_')
        with os.fdopen(fd, 'wb') as f:
            f.write(data)
        return tmp, ''
    except Exception as e:
        return None, f'下载失败：{e}'


def load_source(src):
    """加载来源（链接或本地路径）并解析。

    返回 (records, errors, tmp_path_to_cleanup)
    """
    src = (src or '').strip().strip('"').strip("'")
    if not src:
        return [], ['请给出文件链接或服务器上的文件路径'], None

    tmp = None
    if is_url(src):
        tmp, err = download(src)
        if err:
            return [], [err], None
        path = tmp
    else:
        path = src
        if not os.path.exists(path):
            return [], [f'找不到文件：{path}'], None

    records, errors = parse_file(path)
    return records, errors, tmp


def cleanup(tmp):
    if tmp and os.path.exists(tmp):
        try:
            os.remove(tmp)
        except OSError:
            pass
